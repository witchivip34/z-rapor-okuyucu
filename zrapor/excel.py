"""Kontrol edilmiş Z raporları -> Excel.

Sayfalar: Özet (sade günlük döküm), Detay (gün başına tüm rakamlar ve Kontrol kolonu),
Departmanlar, Ödemeler, Kontrol Listesi. Birden çok firma varsa her firmaya ayrı Özet/Detay.
Sorunlu hücreler kırmızı (hata) / sarı (uyarı) / mavi (düzeltme); açıklama hücre notunda.
"""
from __future__ import annotations

import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .kontrol import _t, _tarih, odeme, odeme_turu

F = "Arial"
BASLIK = PatternFill("solid", fgColor="1F3864")
KIRMIZI = PatternFill("solid", fgColor="FFC7CE")
SARI = PatternFill("solid", fgColor="FFEB9C")
MAVI = PatternFill("solid", fgColor="DDEBF7")
ONEM = {"hata": 0, "uyari": 1, "duzeltme": 2}
DOLGU = {"hata": KIRMIZI, "uyari": SARI, "duzeltme": MAVI}
ISARET = {"hata": "✗ ", "uyari": "! ", "duzeltme": "↻ "}
YAZAR = "z-rapor-okuyucu"
TOPLAM = PatternFill("solid", fgColor="D9E1F2")
INCE = Side(style="thin", color="BFBFBF")
CERCEVE = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)
PARA = '#,##0.00'

# alan adı (kontrol) -> Detay kolon anahtarı
ALAN_KOLON = {
    "tarih": ["tarih"], "z_no": ["z_no"], "net_satis": ["ciro"], "mali_toplam": ["ciro"],
    "odemeler": ["nakit", "kart", "diger"], "banka_dokumu": ["kart"],
    "mali_topkdv": ["kdv"], "kdv_dokumu": ["kdv"], "departmanlar": ["dep"],
    "musteri_fis_adedi": ["musteri"], "toplam_fis_adedi": ["toplam_fis"],
    "fis_iptal_tutar": ["iptal_tutar"], "fis_iptal_adet": ["iptal_adet"], "fis_no": ["fis_no"],
    "mali_fis_adedi": ["mali"], "slip_fis_adedi": ["slip"], "kasa_nakit": ["nakit"], "kum_toplam": ["kum_toplam"], "kum_topkdv": ["kum_kdv"],
}


def _en_agir(sorunlar: list[dict]) -> PatternFill:
    """Sorunlar içindeki en ağır seviyenin dolgusu (hata > uyarı > düzeltme)."""
    return DOLGU[min((s["seviye"] for s in sorunlar), key=ONEM.get)]


def _sayfa_adi(ad: str, kullanilan: set) -> str:
    """Excel'in kabul ettiği, tekil, en çok 28 karakterlik sayfa adı."""
    ad = re.sub(r"[\[\]:*?/\\]", "", ad or "Firma")[:28].strip() or "Firma"
    aday, i = ad, 2
    while aday in kullanilan:
        aday, i = f"{ad[:25]} {i}", i + 1
    kullanilan.add(aday)
    return aday


def _bicim(c, fmt=None, kalin=False, dolgu=None, renk=None, hiza=None) -> None:
    """Hücreye ortak yazı tipi, çerçeve, sayı biçimi ve dolgu uygular."""
    c.font = Font(name=F, bold=kalin, color=renk, size=9)
    c.border = CERCEVE
    if fmt:
        c.number_format = fmt
    if dolgu:
        c.fill = dolgu
    if hiza:
        c.alignment = Alignment(horizontal=hiza, vertical="center", wrap_text=True)


def firma_sayfasi(ws, raporlar: list[dict]) -> None:
    """Detay sayfası: gün başına bir satır; KDV oranı başına toplam/matrah/KDV, kümülatifler, Kontrol."""
    ilk = raporlar[0]
    oranlar = sorted({_t(k["oran"]) for r in raporlar for k in r.get("kdv_dokumu") or []
                      if k.get("oran") is not None})
    kol = [("tarih", "Tarih", 11), ("z_no", "Z No", 7), ("fis_no", "Fiş No", 7),
           ("dep", "Departmanlar", 26)]
    for o in oranlar:
        kol += [(f"t{o:g}", f"%{o:g} Toplam (KDV dahil)", 13), (f"m{o:g}", f"%{o:g} Matrah", 13),
                (f"k{o:g}", f"%{o:g} KDV", 11)]
    kol += [("ciro", "Toplam Ciro (TL)", 13), ("kdv", "Toplam KDV (TL)", 11),
            ("nakit", "Nakit (TL)", 11), ("kart", "Kredi Kartı (TL)", 13),
            ("diger", "Diğer Ödeme (TL)", 11), ("iptal_adet", "İptal Adet", 7),
            ("iptal_tutar", "İptal Tutarı (TL)", 11), ("musteri", "Müşteri Fiş", 8),
            ("mali", "Mali Fiş", 7), ("slip", "Slip Fiş", 7), ("toplam_fis", "Toplam Fiş", 8),
            ("kum_toplam", "Küm. Toplam (TL)", 14), ("kum_kdv", "Küm. KDV (TL)", 13),
            ("eku", "EKÜ No", 6), ("kaynak", "Kaynak Foto", 14), ("kontrol", "Kontrol", 60)]
    idx = {k: i + 1 for i, (k, _, _) in enumerate(kol)}
    toplanan = [k for k, _, _ in kol if k[0] in "tmk" and k[1:2].isdigit()] + \
        ["ciro", "kdv", "nakit", "kart", "diger", "iptal_adet", "iptal_tutar", "musteri", "mali",
         "slip", "toplam_fis"]

    tarihler = [d for d in (_tarih(r.get("tarih")) for r in raporlar) if d]
    donem = f"{min(tarihler):%d.%m.%Y} – {max(tarihler):%d.%m.%Y}" if tarihler else "?"
    ws["A1"] = f"{ilk.get('firma_unvan') or ''} — {ilk.get('firma_sahibi') or ''}".strip(" —")
    ws["A1"].font = Font(name=F, bold=True, size=13, color="1F3864")
    ws["A2"] = (f"VKN: {ilk.get('vkn') or '?'}  ·  VD: {ilk.get('vergi_dairesi') or '?'}  ·  "
                f"Cihaz: {ilk.get('cihaz_sicil_no') or '?'}  ·  Dönem: {donem}  ·  "
                f"{len(raporlar)} Z raporu")
    ws["A3"] = ("Kırmızı: rakamlar birbirini tutmuyor, fişe bakın.  Sarı: olağan dışı ya da tahmini, "
                "kontrol edin.  Mavi: okuma hatası önceki+sonraki Z'den kanıtlanıp düzeltildi.  "
                "Açıklama hücre notunda ve Kontrol kolonunda.  Matrah = Toplam − KDV.")
    for a in ("A2", "A3"):
        ws[a].font = Font(name=F, italic=True, size=9, color="595959")

    H = 5
    for k, ad, gen in kol:
        c = ws.cell(H, idx[k], ad)
        _bicim(c, kalin=True, dolgu=BASLIK, renk="FFFFFF", hiza="center")
        ws.column_dimensions[get_column_letter(idx[k])].width = gen
    ws.row_dimensions[H].height = 30

    satir = H
    for r in raporlar:
        satir += 1
        kd = {_t(k["oran"]): k for k in r.get("kdv_dokumu") or [] if k.get("oran") is not None}
        dep = "; ".join(f"{d.get('ad')} %{_t(d.get('kdv_orani')):g}: "
                        f"{_t(d.get('tutar')):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                        for d in r.get("departmanlar") or [])
        d = _tarih(r.get("tarih"))
        deger = {"tarih": d or r.get("tarih"), "z_no": r.get("z_no"), "fis_no": r.get("fis_no"),
                 "dep": dep, "ciro": r.get("net_satis"), "kdv": r.get("mali_topkdv"),
                 "nakit": odeme(r, "nakit") or 0, "kart": odeme(r, "kart") or 0,
                 "diger": odeme(r, "diger") or 0, "iptal_adet": r.get("fis_iptal_adet") or 0,
                 "iptal_tutar": r.get("fis_iptal_tutar") or 0,
                 "musteri": r.get("musteri_fis_adedi"), "mali": r.get("mali_fis_adedi"),
                 "slip": r.get("slip_fis_adedi"), "toplam_fis": r.get("toplam_fis_adedi"),
                 "kum_toplam": r.get("kum_toplam"), "kum_kdv": r.get("kum_topkdv"),
                 "eku": r.get("eku_no"), "kaynak": r.get("_kaynak")}
        for o in oranlar:
            k = kd.get(o)
            deger[f"t{o:g}"] = k["toplam"] if k else 0
            deger[f"k{o:g}"] = k["kdv"] if k else 0
            deger[f"m{o:g}"] = (f"={get_column_letter(idx[f't{o:g}'])}{satir}"
                                f"-{get_column_letter(idx[f'k{o:g}'])}{satir}")
        sorunlar = r.get("_sorunlar") or []
        deger["kontrol"] = "✓" if not sorunlar else " | ".join(
            ISARET[s["seviye"]] + s["mesaj"] for s in sorunlar)

        isaret: dict[str, list] = {}
        for s in sorunlar:
            alan = s["alan"]
            hedef = ALAN_KOLON.get(alan) or ([f"k{alan[4:]}"] if alan.startswith("kdv_") else [])
            for h in hedef:
                if h in idx:
                    isaret.setdefault(h, []).append(s)

        for k, _, _ in kol:
            c = ws.cell(satir, idx[k], deger.get(k))
            fmt = "DD.MM.YYYY" if k == "tarih" else (PARA if (k in toplanan and k not in
                  ("iptal_adet", "musteri", "mali", "slip", "toplam_fis")) or k.startswith("kum") else None)
            dolgu = None
            if k in isaret:
                dolgu = _en_agir(isaret[k])
                c.comment = Comment("\n".join(s["mesaj"] for s in isaret[k]), YAZAR)
            if k == "kontrol" and sorunlar:
                dolgu = _en_agir(sorunlar)
            _bicim(c, fmt=fmt, dolgu=dolgu,
                   hiza="left" if k in ("dep", "kontrol", "kaynak") else "center"
                   if k in ("tarih", "z_no", "fis_no", "eku") else None)

    son = satir + 1
    ws.cell(son, 1, "GENEL TOPLAM")
    for k, _, _ in kol:
        c = ws.cell(son, idx[k])
        if k in toplanan:
            harf = get_column_letter(idx[k])
            c.value = f"=SUM({harf}{H + 1}:{harf}{satir})"
        _bicim(c, fmt=PARA if k in toplanan and k not in
               ("iptal_adet", "musteri", "mali", "slip", "toplam_fis") else None,
               kalin=True, dolgu=TOPLAM)
    ws.freeze_panes = ws.cell(H + 1, idx["dep"])


AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim",
         "Kasım", "Aralık"]
TL_FMT = '#,##0.00 "TL"'
# alan (kontrol) -> Özet kolonu; karşılığı olmayan hata Tarih hücresine not olarak düşer
OZET_ALAN = {"tarih": "tarih", "z_no": "z", "departmanlar": "dep", "net_satis": "ciro", "mali_toplam": "ciro",
             "mali_topkdv": "kdv", "kdv_dokumu": "kdv", "odemeler": "nakit kart diger",
             "banka_dokumu": "kart", "fis_iptal_tutar": "iptal", "fis_iptal_adet": "iptal"}


def _donem_yazisi(tarihler) -> str:
    """'01 - 30 Eylül 2026' gibi dönem yazısı."""
    if not tarihler:
        return ""
    a, b = min(tarihler), max(tarihler)
    if (a.year, a.month) == (b.year, b.month):
        return f"{a.day:02d} - {b.day:02d} {AYLAR[a.month - 1]} {a.year}"
    return f"{a:%d.%m.%Y} - {b:%d.%m.%Y}"


def ozet_sayfasi(ws, raporlar: list[dict], detay_adi: str):
    """Sade günlük döküm: Tarih, Z No, Departman, KDV oranı/tutarı, ciro, nakit, kart, iptal (TL)."""
    ilk = raporlar[0]
    diger_var = any(odeme(r, "diger") for r in raporlar)
    kol = [("tarih", "Tarih", 12), ("z", "Z No", 7), ("dep", "Departman", 16), ("oran", "KDV Oranı", 9),
           ("kdv", "KDV Tutarı", 14), ("ciro", "Toplam Ciro (TL)", 16), ("nakit", "Nakit (TL)", 14),
           ("kart", "Kredi Kartı (TL)", 16)] + ([("diger", "Diğer Ödeme (TL)", 14)] if diger_var else []) + \
          [("iptal", "İptal Tutarı (TL)", 14)]
    idx = {k: i + 1 for i, (k, _, _) in enumerate(kol)}
    son_kol = get_column_letter(len(kol))

    tarihler = [d for d in (_tarih(r.get("tarih")) for r in raporlar) if d]
    ws["A1"] = f"{ilk.get('firma_unvan') or ''} - {ilk.get('firma_sahibi') or ''}".strip(" -")
    ws["A1"].font = Font(name=F, bold=True, size=13, color="1F3864")
    ws["A2"] = f"Z Raporları Günlük Satış ve İşlem Dökümü ({_donem_yazisi(tarihler)})"
    ws["A2"].font = Font(name=F, italic=True, size=9, color="595959")
    ws.merge_cells(f"A1:{son_kol}1")
    ws.merge_cells(f"A2:{son_kol}2")

    H = 4
    for k, ad, gen in kol:
        c = ws.cell(H, idx[k], ad)
        _bicim(c, kalin=True, dolgu=BASLIK, renk="FFFFFF", hiza="center")
        ws.column_dimensions[get_column_letter(idx[k])].width = gen
    ws.row_dimensions[H].height = 30

    satir = H
    for r in raporlar:
        satir += 1
        oranlar = sorted({_t(k["oran"]) for k in r.get("kdv_dokumu") or [] if k.get("oran") is not None}
                         or {_t(d["kdv_orani"]) for d in r.get("departmanlar") or [] if d.get("kdv_orani") is not None})
        adlar = []
        for d in r.get("departmanlar") or []:
            if d.get("ad") and d["ad"] not in adlar:
                adlar.append(d["ad"])
        try:
            z = int(str(r.get("z_no")))
        except (TypeError, ValueError):
            z = r.get("z_no")
        deger = {"tarih": _tarih(r.get("tarih")) or r.get("tarih"), "z": z, "dep": ", ".join(adlar),
                 "oran": oranlar[0] / 100 if len(oranlar) == 1 else " / ".join(f"%{o:g}" for o in oranlar),
                 "kdv": r.get("mali_topkdv"), "ciro": r.get("net_satis"),
                 "nakit": odeme(r, "nakit") or 0, "kart": odeme(r, "kart") or 0,
                 "diger": odeme(r, "diger") or 0, "iptal": r.get("fis_iptal_tutar") or 0}

        isaret: dict[str, list] = {}
        for x in r.get("_sorunlar") or []:
            hedef = OZET_ALAN.get(x["alan"]) or ("kdv" if x["alan"].startswith("kdv_") else None)
            if hedef is None and x["seviye"] == "hata":
                hedef = "tarih"  # ayrintisi Detay sayfasinda
            for h in (hedef or "").split():
                if h in idx:
                    isaret.setdefault(h, []).append(x)

        for k, _, _ in kol:
            c = ws.cell(satir, idx[k], deger.get(k))
            fmt = ("DD.MM.YYYY" if k == "tarih" else "0%" if k == "oran" and isinstance(deger[k], float)
                   else TL_FMT if k in ("kdv", "ciro", "nakit", "kart", "diger", "iptal") else None)
            dolgu = None
            if k in isaret:
                dolgu = _en_agir(isaret[k])
                c.comment = Comment("\n".join(x["mesaj"] for x in isaret[k])
                                    + f"\n(Ayrıntı: '{detay_adi}' sayfası)", YAZAR)
            _bicim(c, fmt=fmt, dolgu=dolgu, hiza="center" if k in ("tarih", "z", "oran") else
                   "left" if k == "dep" else None)

    son = satir + 1
    ws.cell(son, 1, "GENEL TOPLAM")
    for k, _, _ in kol:
        c = ws.cell(son, idx[k])
        if k in ("kdv", "ciro", "nakit", "kart", "diger", "iptal"):
            h = get_column_letter(idx[k])
            c.value = f"=SUM({h}{H + 1}:{h}{satir})"
        _bicim(c, fmt=TL_FMT, kalin=True, dolgu=TOPLAM)
    ws.cell(son + 2, 1, "Renkli hücre: kırmızı = fişe bakın, sarı = kontrol edin, mavi = okuma hatası "
            f"zincirden düzeltildi. Notu görmek için hücrenin üstüne gelin; ayrıntılı döküm '{detay_adi}' sayfasında.")
    ws.cell(son + 2, 1).font = Font(name=F, italic=True, size=8, color="595959")
    ws.freeze_panes = ws.cell(H + 1, 1)


def yaz(firmalar: dict[str, list[dict]], cikti) -> None:
    """Kontrol edilmiş firmaları (kontrol.kontrol_et çıktısı) tek Excel dosyasına yazar."""
    wb = Workbook()
    wb.remove(wb.active)
    kullanilan: set = set()
    tek = len(firmalar) == 1
    for _, raporlar in firmalar.items():
        firma = (raporlar[0].get("firma_unvan") or "Firma")[:20]
        ozet_adi = _sayfa_adi("Özet" if tek else f"Özet {firma}", kullanilan)
        detay_adi = _sayfa_adi("Detay" if tek else f"Detay {firma}", kullanilan)
        ozet_ws = wb.create_sheet(ozet_adi)
        firma_sayfasi(wb.create_sheet(detay_adi), raporlar)
        ozet_sayfasi(ozet_ws, raporlar, detay_adi)

    def tablo(ad, basliklar, satirlar, genislik):
        ws = wb.create_sheet(ad)
        for i, b in enumerate(basliklar, 1):
            _bicim(ws.cell(1, i, b), kalin=True, dolgu=BASLIK, renk="FFFFFF", hiza="center")
            ws.column_dimensions[get_column_letter(i)].width = genislik[i - 1]
        for ri, s in enumerate(satirlar, 2):
            for ci, v in enumerate(s, 1):
                c = ws.cell(ri, ci, v)
                _bicim(c, fmt="DD.MM.YYYY" if ci == 2 else PARA if isinstance(v, float) else None)
        ws.freeze_panes = "A2"
        return ws

    hepsi = [(r, r.get("firma_unvan")) for rs in firmalar.values() for r in rs]
    tablo("Departmanlar", ["Firma", "Tarih", "Z No", "Departman", "KDV %", "Adet", "Tutar (TL)"],
          [[f, _tarih(r.get("tarih")), r.get("z_no"), d.get("ad"), d.get("kdv_orani"),
            d.get("adet"), _t(d.get("tutar"))] for r, f in hepsi for d in r.get("departmanlar") or []],
          [22, 11, 7, 16, 7, 7, 13])
    tablo("Ödemeler", ["Firma", "Tarih", "Z No", "Tür", "Ayrıntı", "Adet", "Tutar (TL)"],
          [[f, _tarih(r.get("tarih")), r.get("z_no"), odeme_turu(o.get("tur")), o.get("tur"),
            o.get("adet"), _t(o.get("tutar"))] for r, f in hepsi for o in r.get("odemeler") or []] +
          [[f, _tarih(r.get("tarih")), r.get("z_no"), "banka", b.get("banka"), b.get("adet"),
            _t(b.get("tutar"))] for r, f in hepsi for b in r.get("banka_dokumu") or []],
          [22, 11, 7, 8, 16, 7, 13])
    kl = [[f, _tarih(r.get("tarih")), r.get("z_no"), s["seviye"].upper(), s["alan"], s["mesaj"],
           r.get("_kaynak")] for r, f in hepsi for s in r.get("_sorunlar") or []]
    ws = tablo("Kontrol Listesi", ["Firma", "Tarih", "Z No", "Seviye", "Alan", "Açıklama", "Foto"],
               kl or [["Tüm kontroller geçti", None, None, None, None, None, None]],
               [22, 11, 7, 8, 16, 70, 16])
    for row in ws.iter_rows(min_row=2):
        if (row[3].value or "").lower() in DOLGU:
            row[3].fill = DOLGU[row[3].value.lower()]
    ws.cell(len(kl) + 3, 1, f"Oluşturma: {datetime.now():%d.%m.%Y %H:%M} · z-rapor-okuyucu. "
            "Rakamlar fotoğraftan yapay zekâ ile okunmuştur; işaretli hücreler fişle karşılaştırılmalıdır.")
    wb.save(cikti)

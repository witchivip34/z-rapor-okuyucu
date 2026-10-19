"""Kontrol edilmiş Z raporları -> Luca içe aktarım dosyaları.

Bilanço: Luca "Fiş Aktarım Şablon" (.xlsx, `sablonlar/bilanco_fis_aktarim_sablon.xlsx`).
    Her Z bir mahsup fişi: KDV oranı başına 600 (matrah, alacak) + 391 (KDV, alacak);
    kart -> 108 POS (borç), nakit -> 100 Kasa (borç). Belge türü ZR, evrak no = Z no.
İşletme defteri: `sablonlar/isletme_hizli_fis_aktarim.csv` (cp1254, ';', 36 kolon, ondalık virgül).
    Z başına KDV oranı başına bir satır.

Hesap kodları: varsayılanlar aşağıda (tekdüzen hesap planı). Firma (VKN) başına geçersiz kılma
kullanıcının kendi `hesaplar.json` dosyasından okunur; biçim için `ornekler/hesaplar.ornek.json`.

Not: Luca'ya gerçek içe aktarım henüz doğrulanmadı; sütun değerleri örnek kayıtlardan çıkarımdır.
"""
from __future__ import annotations

import json
import re
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook

from .kontrol import _t, _tarih, odeme

SABLONLAR = Path(__file__).parent / "sablonlar"
SABLON_BILANCO = SABLONLAR / "bilanco_fis_aktarim_sablon.xlsx"
SABLON_ISLETME = SABLONLAR / "isletme_hizli_fis_aktarim.csv"

# Varsayılan alt hesaplar. Firmanın hesap planı farklıysa hesaplar.json ile değiştirilir.
VARSAYILAN: dict = {
    "kasa": "100.01.001",
    "pos": "108.01.001",
    "diger_odeme": "108.01.001",
    "satis": {"1": "600.01.001", "10": "600.01.002", "20": "600.01.003"},
    "kdv": {"1": "391.01.001", "10": "391.01.002", "20": "391.01.003"},
    "belge_turu": "ZR",
    "aciklama": "Z RAPORU",
    "fis_no": "z_no",  # "z_no": fiş no = Z no | "bos": Luca numaralasın
}
# Firma dosyasında verilmediği sürece "kontrol edilmemiş kod" sayılmayan anahtarlar
DOGRULANMIS = {"kasa", "pos", "satis.10", "kdv.10", "belge_turu"}

# İşletme defteri kod kolonları (örnek bir Luca Z raporu kaydından)
ISLETME_VARSAYILAN: dict = {
    "islem": "Gelir", "kategori": "Defter Fişleri", "belge_turu": "Alış", "kod": "1100",
    "belge_turu_db": "Z Raporu", "satis_turu": "Normal Satışlar", "kayit_alt_turu": "Mal Satışı",
    "tevkifat": "Yok", "faaliyet_kodu": "",
}


# --- hesap kodları ---------------------------------------------------------------------------
def hesaplar_oku(yol: Path | str | None) -> dict:
    """Kullanıcının hesaplar.json dosyasını okur ({VKN: {...}}); yol yoksa boş sözlük."""
    if yol and Path(yol).exists():
        return json.loads(Path(yol).read_text(encoding="utf-8"))
    return {}


def hesaplar_yaz(tum: dict, yol: Path | str) -> None:
    """Hesap kodlarını JSON olarak kaydeder."""
    p = Path(yol)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(tum, ensure_ascii=False, indent=1), encoding="utf-8")


def vkn_anahtari(vkn: str | None) -> str:
    """VKN'den yalnız rakamlar (en çok 11): hesaplar.json anahtarı."""
    return re.sub(r"\D", "", vkn or "")[:11]


def firma_hesaplari(vkn: str | None, tum: dict | None = None) -> dict:
    """Varsayılan kodların üstüne firmanın kendi kodlarını bindirir."""
    h = deepcopy(VARSAYILAN)
    for k, v in (tum or {}).get(vkn_anahtari(vkn), {}).items():
        if isinstance(v, dict):
            h.setdefault(k, {}).update(v)
        elif v not in (None, ""):
            h[k] = v
    return h


def _oran_anahtari(o) -> str:
    return f"{_t(o):g}"


def aktarilamaz(r: dict) -> str | None:
    """Luca'ya gitmemesi gereken Z için nedeni döndürür: para alanlarında çözülmemiş hata ya da eksik rakam."""
    para = {"net_satis", "mali_toplam", "mali_topkdv", "kdv_dokumu", "odemeler", "tarih"}
    # "Z 131 fotografta yok" bu Z'nin değil, öncekinin eksikliği: bu fişi engellemez
    hatalar = [s["mesaj"] for s in r.get("_sorunlar") or []
               if s["seviye"] == "hata" and (s["alan"] in para or s["alan"].startswith("kdv_")
                                             or (s["alan"] == "z_no" and "fotografta yok" not in s["mesaj"]))]
    if hatalar:
        return "; ".join(hatalar)
    if not _tarih(r.get("tarih")) or r.get("net_satis") is None:
        return "tarih ya da net satis yok"
    kd = r.get("kdv_dokumu") or []
    if not kd or any(k.get("toplam") is None or k.get("kdv") is None or k.get("oran") is None for k in kd):
        return "vergi dokumu eksik"
    return None


def _aktarilmayan(r: dict, neden: str) -> dict:
    return {"firma": r.get("firma_unvan"), "tarih": r.get("tarih"), "z_no": r.get("z_no"), "neden": neden}


# --- bilanço -----------------------------------------------------------------------------------
def fis_satirlari(r: dict, h: dict) -> list[dict]:
    """Bir Z raporunun dengeli (borç = alacak) mahsup satırları."""
    aciklama = f"{h['aciklama']} {r.get('cihaz_sicil_no') or ''}".strip()
    satirlar = []
    for k in sorted(r["kdv_dokumu"], key=lambda k: _t(k["oran"])):
        o = _oran_anahtari(k["oran"])
        matrah = round(_t(k["toplam"]) - _t(k["kdv"]), 2)
        if matrah:
            satirlar.append({"hesap": h["satis"].get(o) or f"600.?? (%{o})", "borc": 0, "alacak": matrah,
                             "detay": aciklama})
        if _t(k["kdv"]):
            satirlar.append({"hesap": h["kdv"].get(o) or f"391.?? (%{o})", "borc": 0,
                             "alacak": round(_t(k["kdv"]), 2), "detay": aciklama})
    for tur, hesap in (("kart", h["pos"]), ("nakit", h["kasa"]), ("diger", h["diger_odeme"])):
        tutar = round(odeme(r, tur) or 0, 2)
        if tutar:
            satirlar.append({"hesap": hesap, "borc": tutar, "alacak": 0, "detay": aciklama})
    # kuruş dengesi: ödemeler ile vergi dökümü arasında yuvarlama farkı kalırsa matraha yaz
    fark = round(sum(s["borc"] for s in satirlar) - sum(s["alacak"] for s in satirlar), 2)
    if fark and abs(fark) <= 0.05:
        for s in satirlar:
            if s["alacak"] and str(s["hesap"]).startswith("600"):
                s["alacak"] = round(s["alacak"] + fark, 2)
                break
    return satirlar


def _tahmini_kod(s: dict, h: dict, ozel: dict) -> str | None:
    """Satırın hesap kodu ne doğrulanmış varsayılan ne de firmanın kendi kodu ise etiketini döndürür."""
    anahtar = next((f"{a}.{o}" for a in ("satis", "kdv") for o, kod in h[a].items() if kod == s["hesap"]), None)
    if anahtar and anahtar not in DOGRULANMIS and anahtar.split(".")[1] not in (ozel.get(anahtar.split(".")[0]) or {}):
        return f"{s['hesap']} ({anahtar.replace('satis.', 'satış %').replace('kdv.', 'KDV %')})"
    return None


def bilanco_yaz(firmalar: dict[str, list[dict]], cikti: Path, hesaplar: dict | None = None) -> dict:
    """Luca bilanço fiş aktarım dosyası yazar.

    Döner: fiş/satır sayısı, aktarılmayan Z'ler (nedeniyle), kontrol edilmemiş (tahmini) kodlar.
    """
    wb = load_workbook(SABLON_BILANCO)
    ws = wb.worksheets[0]
    bicim = {c.column: c.number_format for c in ws[1]}
    satir_no, fis, aktarilmayan, tahmini = 2, 0, [], set()
    for raporlar in firmalar.values():
        vkn = raporlar[0].get("vkn")
        h = firma_hesaplari(vkn, hesaplar)
        ozel = (hesaplar or {}).get(vkn_anahtari(vkn), {})
        for r in raporlar:
            neden = aktarilamaz(r)
            if neden:
                aktarilmayan.append(_aktarilmayan(r, neden))
                continue
            fis += 1
            tarih = _tarih(r["tarih"])
            z = str(r.get("z_no") or "")
            fis_no = (z.lstrip("0") or z) if h.get("fis_no") == "z_no" else None
            aciklama = f"{h['aciklama']} {r.get('cihaz_sicil_no') or ''}".strip()
            for s in fis_satirlari(r, h):
                etiket = _tahmini_kod(s, h, ozel)
                if etiket:
                    tahmini.add(etiket)
                degerler = [fis_no, tarih, aciklama, s["hesap"], z.lstrip("0") or z, tarih,
                            s["detay"], s["borc"], s["alacak"], 0, h["belge_turu"], None, None, None, None]
                for col, v in enumerate(degerler, 1):
                    c = ws.cell(satir_no, col, v)
                    c.number_format = bicim.get(col, "General")
                satir_no += 1
    wb.save(cikti)
    return {"fis": fis, "satir": satir_no - 2, "aktarilmayan": aktarilmayan, "tahmini_kodlar": sorted(tahmini)}


# --- işletme defteri ---------------------------------------------------------------------------
def _tl_csv(x) -> str:
    return f"{_t(x):.2f}".replace(".", ",")


def isletme_basliklari() -> list[str]:
    """Şablondaki 36 kolon başlığı."""
    return SABLON_ISLETME.read_bytes().decode("cp1254").strip().split(";")


def _isletme_satirlari(r: dict, h: dict) -> list[dict]:
    """Bir Z'nin KDV oranı başına satırları; kartla tahsilat oranlara toplam payıyla dağıtılır."""
    tarih = _tarih(r["tarih"]).strftime("%d/%m/%Y")
    z = str(r.get("z_no") or "").lstrip("0") or str(r.get("z_no") or "")
    kd = sorted([k for k in r["kdv_dokumu"] if _t(k["toplam"])], key=lambda k: _t(k["oran"]))
    kart = round(odeme(r, "kart") or 0, 2)
    toplam = sum(_t(k["toplam"]) for k in kd) or 1
    dagitilan, satirlar = 0.0, []
    for i, k in enumerate(kd):
        # kuruş farkı son satıra
        pay = round(kart - dagitilan, 2) if i == len(kd) - 1 else round(kart * _t(k["toplam"]) / toplam, 2)
        dagitilan += pay
        matrah = round(_t(k["toplam"]) - _t(k["kdv"]), 2)
        satirlar.append({
            "İŞLEM": h["islem"], "KATEGORİ": h["kategori"], "BELGE TURU": h["belge_turu"],
            "EVRAK TARİHİ": tarih, "KAYIT TARİHİ": tarih, "EVRAK NO": z, "KOD": h["kod"],
            "BELGE TÜRÜ(DB)": h["belge_turu_db"], "ALIŞ/SATIŞ TÜRÜ": h["satis_turu"],
            "KAYIT ALT TÜRÜ": h["kayit_alt_turu"], "AÇIKLAMA": "Z RAPORU",
            "MİKTAR": "0,00", "B.FİYAT": "0,00", "TUTAR": _tl_csv(matrah), "TEVKİFAT": h["tevkifat"],
            "KDV ORANI": f"{_t(k['oran']):g}", "ÖZEL MATRAH İŞLEM BEDELİ": "0,00",
            "MATRAHTAN DÜŞÜLECEK TUTAR": "0,00", "MATRAHA DAHİL OLMAYAN BEDEL": "0,00",
            "KDV TUTARI": _tl_csv(k["kdv"]), "TOPLAM TUTAR": _tl_csv(k["toplam"]),
            "KREDİLİ TUTAR": _tl_csv(pay), "STOPAJ TUTARI": "0,00", "FAALIYET KODU": h["faaliyet_kodu"]})
    return satirlar


def isletme_yaz(firmalar: dict[str, list[dict]], cikti: Path, hesaplar: dict | None = None) -> dict:
    """Luca işletme defteri hızlı fiş aktarım CSV'si yazar (cp1254, ';')."""
    baslik = isletme_basliklari()
    satirlar, fis, aktarilmayan, eksik_faaliyet = [], 0, [], set()
    for raporlar in firmalar.values():
        vkn = raporlar[0].get("vkn")
        ozel = ((hesaplar or {}).get(vkn_anahtari(vkn), {}).get("isletme") or {})
        h = ISLETME_VARSAYILAN | {k: v for k, v in ozel.items() if v}
        if not h["faaliyet_kodu"]:
            eksik_faaliyet.add(raporlar[0].get("firma_unvan") or vkn or "?")
        for r in raporlar:
            neden = aktarilamaz(r)
            if neden:
                aktarilmayan.append(_aktarilmayan(r, neden))
                continue
            fis += 1
            for deger in _isletme_satirlari(r, h):
                satirlar.append(";".join(str(deger.get(b, "")).replace(";", ",") for b in baslik))
    metin = ";".join(baslik) + "\r\n" + "".join(s + "\r\n" for s in satirlar)
    Path(cikti).write_bytes(metin.encode("cp1254", errors="replace"))
    return {"fis": fis, "satir": len(satirlar), "aktarilmayan": aktarilmayan,
            "eksik_faaliyet": sorted(eksik_faaliyet)}

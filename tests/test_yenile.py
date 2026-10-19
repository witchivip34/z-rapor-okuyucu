"""`z-rapor yenile`: model çağrısı yapmadan önbellekten Excel."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import ORNEKLER

from zrapor import okuma
from zrapor.cli import main


def _onbellege_yaz(klasor: Path, okuma_sayisi: int = 2) -> None:
    """Örnek okumayı `oku_ham`ın yazdığı biçimde (ham model çıktısı + _kaynak) önbelleğe koyar."""
    raporlar = json.loads((ORNEKLER / "okuma_ornek.json").read_text(encoding="utf-8"))["raporlar"]
    raporlar = [{k: v for k, v in r.items() if not k.startswith("_")} for r in raporlar]
    klasor.mkdir(parents=True, exist_ok=True)
    for n in range(1, okuma_sayisi + 1):
        ad = f"fis_01_0123456789ab_sonnet_api{'' if n == 1 else f'_o{n}'}.json"
        (klasor / ad).write_text(json.dumps({"raporlar": raporlar, "z_raporu_olmayan": [],
                                             "_kaynak": "fis_01.jpg"}), encoding="utf-8")


@pytest.fixture(autouse=True)
def model_cagrisi_yasak(monkeypatch):
    def yasak(*a, **k):
        raise AssertionError("yenile model cagirmamali")
    monkeypatch.setattr(okuma, "_api", yasak)
    monkeypatch.setattr(okuma, "_gemini", yasak)
    monkeypatch.setattr(okuma, "_cli", yasak)


def test_yenile_cikti_dosyasi(tmp_path: Path, capsys):
    _onbellege_yaz(tmp_path / "_okumalar")
    cikti = tmp_path / "eylul.xlsx"
    main(["yenile", str(cikti)])
    out = capsys.readouterr().out
    assert "1 fotograf, 2 okuma" in out and "5 Z · 1 hata" in out
    assert "Detay" in load_workbook(cikti).sheetnames
    firmalar = json.loads((tmp_path / "raporlar.json").read_text(encoding="utf-8"))
    assert {r["_kaynak"] for r in firmalar["1234567890"]} == {"fis_01.jpg"}


def test_yenile_onbellek_yoksa_hata(tmp_path: Path):
    with pytest.raises(FileNotFoundError):
        main(["yenile", str(tmp_path / "yok.xlsx")])


def test_yenile_web_isi(tmp_path: Path, capsys):
    is_id = "20260919-101500-abc123"
    klasor = tmp_path / "isler" / is_id
    _onbellege_yaz(klasor / "cikti" / "_okumalar")
    durum = {"id": is_id, "durum": "bitti", "tur": "z", "excel": "ORNEK_LOKANTA.xlsx", "okuma": 2,
             "mesajlar": [], "ozet": []}
    (klasor / "durum.json").write_text(json.dumps(durum), encoding="utf-8")
    main(["yenile", "--veri", str(tmp_path), "--is", is_id])
    assert (klasor / "cikti" / "ORNEK_LOKANTA.xlsx").exists()
    yeni = json.loads((klasor / "durum.json").read_text(encoding="utf-8"))
    assert yeni["ozet"][0]["vkn"] == "1234567890" and yeni["ozet"][0]["z"] == 5
    assert "ORNEK_LOKANTA.xlsx" in capsys.readouterr().out

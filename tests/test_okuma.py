"""Okuma hattı: parçalama, önbellek ve uçtan uca işleme (API çağrısı sahte)."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from conftest import ORNEKLER

from zrapor import okuma
from zrapor.isleme import isle


def _sahte_api(cagrilar: list):
    rapor = json.loads((ORNEKLER / "okuma_ornek.json").read_text(encoding="utf-8"))["raporlar"][0]
    rapor = {k: v for k, v in rapor.items() if not k.startswith("_")}

    def api(gorseller, model, sema, talimat):
        cagrilar.append((len(gorseller), model))
        return {"raporlar": [rapor], "z_raporu_olmayan": [{"tur": "POS slibi", "satici": None}]}
    return api


def test_parcala_tam_ve_dokuz_parca(tmp_path: Path):
    yollar = okuma.parcala(ORNEKLER / "z_ornek.png", tmp_path / "p")
    assert len(yollar) == 10 and all(p.exists() for p in yollar)


def test_oku_onbellek_ve_z_degil(tmp_path: Path, monkeypatch):
    cagrilar: list = []
    monkeypatch.setattr(okuma, "_api", _sahte_api(cagrilar))
    foto = ORNEKLER / "z_ornek.png"
    r1 = okuma.oku(foto, tmp_path, model="sonnet", arka_uc="api")
    r2 = okuma.oku(foto, tmp_path, model="sonnet", arka_uc="api")
    assert len(cagrilar) == 1  # ikinci okuma önbellekten
    assert cagrilar[0] == (10, "claude-sonnet-5-5")
    assert r1 == r2 and r1[0]["_kaynak"] == "z_ornek.png" and "_z_degil" in r1[1]


def test_isle_iki_okuma_excel(tmp_path: Path, monkeypatch):
    cagrilar: list = []
    monkeypatch.setattr(okuma, "_api", _sahte_api(cagrilar))
    ozet = isle([ORNEKLER / "z_ornek.png"], tmp_path / "z.xlsx", okuma=2, paralel=2, bildir=lambda m: None)
    assert len(cagrilar) == 2  # iki bağımsız okuma, ayrı önbellek kaydı
    assert ozet[0]["vkn"] == "1234567890" and ozet[0]["z"] == 1 and ozet[0]["hata"] == 0
    assert "Özet" in load_workbook(tmp_path / "z.xlsx").sheetnames

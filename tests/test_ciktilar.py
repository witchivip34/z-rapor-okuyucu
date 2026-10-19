"""Excel, Luca ve CLI uçtan uca (API'siz)."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from conftest import ORNEKLER

from zrapor import luca
from zrapor.cli import main
from zrapor.excel import KIRMIZI, MAVI, yaz
from zrapor.kontrol import kontrol_et


def _renk(c) -> str:
    return (c.fill.fgColor.rgb or "")[-6:]


def test_excel_sayfalar_ve_renkler(ornek_raporlar, tmp_path: Path):
    cikti = tmp_path / "z.xlsx"
    yaz(kontrol_et(ornek_raporlar), cikti)
    wb = load_workbook(cikti)
    assert wb.sheetnames == ["Özet", "Detay", "Departmanlar", "Ödemeler", "Kontrol Listesi"]
    detay = wb["Detay"]
    basliklar = {c.value: c.column for c in detay[5]}
    satir = {detay.cell(r, basliklar["Z No"]).value: r for r in range(6, 11)}
    # Z 0103: ciro hücresi kırmızı; Z 0104: kümülatif toplam hücresi mavi
    assert _renk(detay.cell(satir["0103"], basliklar["Toplam Ciro (TL)"])) == KIRMIZI.fgColor.rgb[-6:]
    assert _renk(detay.cell(satir["0104"], basliklar["Küm. Toplam (TL)"])) == MAVI.fgColor.rgb[-6:]
    assert detay.cell(satir["0101"], basliklar["Kontrol"]).value == "✓"
    seviyeler = [r[3].value for r in wb["Kontrol Listesi"].iter_rows(min_row=2) if r[3].value]
    assert "HATA" in seviyeler and "DUZELTME" in seviyeler


def test_luca_bilanco_dengeli(ornek_raporlar, tmp_path: Path):
    firmalar = kontrol_et(ornek_raporlar)
    hesaplar = luca.hesaplar_oku(ORNEKLER / "hesaplar.ornek.json")
    cikti = tmp_path / "bilanco.xlsx"
    o = luca.bilanco_yaz(firmalar, cikti, hesaplar)
    assert o["fis"] == 4 and len(o["aktarilmayan"]) == 1  # kırmızı Z 0103 aktarılmaz
    assert o["tahmini_kodlar"] == []  # örnek hesaplar dosyası %20 kodlarını veriyor
    ws = load_workbook(cikti).worksheets[0]
    fisler: dict = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        b = fisler.setdefault(row[4], [0.0, 0.0])  # Evrak No = Z no
        b[0] += row[7] or 0
        b[1] += row[8] or 0
    assert len(fisler) == 4
    for borc, alacak in fisler.values():
        assert round(borc, 2) == round(alacak, 2)


def test_luca_isletme_csv_36_kolon(ornek_raporlar, tmp_path: Path):
    cikti = tmp_path / "isletme.csv"
    o = luca.isletme_yaz(kontrol_et(ornek_raporlar), cikti,
                         luca.hesaplar_oku(ORNEKLER / "hesaplar.ornek.json"))
    satirlar = cikti.read_bytes().decode("cp1254").strip().split("\r\n")
    assert o["satir"] == 8 and len(satirlar) == 9  # başlık + 4 Z × 2 oran
    assert all(len(s.split(";")) == 36 for s in satirlar)
    assert o["eksik_faaliyet"] == []


def test_sablonlarda_veri_yok():
    csv = luca.SABLON_ISLETME.read_bytes().decode("cp1254").strip().split("\r\n")
    assert len(csv) == 1
    ws = load_workbook(luca.SABLON_BILANCO).worksheets[0]
    assert ws.max_row == 1


def test_cli_dogrula_ve_luca(tmp_path: Path, capsys):
    cikti = tmp_path / "ornek.xlsx"
    main(["dogrula", str(ORNEKLER / "okuma_ornek.json"), "-o", str(cikti)])
    assert cikti.exists()
    raporlar = tmp_path / "raporlar.json"
    firmalar = json.loads(raporlar.read_text(encoding="utf-8"))
    assert list(firmalar) == ["1234567890"]
    assert "1 hata" in capsys.readouterr().out
    main(["luca", str(raporlar), "--tur", "bilanco", "--hesaplar",
          str(ORNEKLER / "hesaplar.ornek.json"), "-o", str(tmp_path / "b.xlsx")])
    assert "4 fiş" in capsys.readouterr().out
